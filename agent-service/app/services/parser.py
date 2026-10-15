"""文档解析与切分。

- ``parse_bytes``：按扩展名/魔数解析 pdf（pypdf）、docx（python-docx）、txt/md（纯文本）。
- ``split_text``：LangChain RecursiveCharacterTextSplitter，分隔符优先中文段落/句读，
  chunk_size / chunk_overlap 来自 kb 服务配置（t_kb.chunk_size / chunk_overlap）。
"""

from __future__ import annotations

import io
import logging
import zipfile

logger = logging.getLogger(__name__)

_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".text", ".log", ""}
_DOCX_SUFFIXES = {".docx"}
_PDF_SUFFIXES = {".pdf"}


def parse_bytes(filename: str, data: bytes) -> str:
    """按扩展名解析文件字节为纯文本；扩展名缺失/不可信时按魔数嗅探兜底。"""
    fname = (filename or "").strip().lower()
    suffix = fname[fname.rfind(".") :] if "." in fname else ""

    if suffix in _PDF_SUFFIXES or data[:4] == b"%PDF":
        text = _parse_pdf(data)
    elif suffix in _DOCX_SUFFIXES or (data[:2] == b"PK" and _looks_docx(data)):
        text = _parse_docx(data)
    else:
        text = _decode_text(data)
    return _normalize(text)


def _normalize(text: str) -> str:
    """PDF/DOCX 提取产物的空白清洗。

    PDF 提取常在汉字之间/行尾插入多余空格与换行，字符 bigram embedding 会因此
    匹配不上连续词（查询「教育背景」vs 块内「教 育 背 景」零重叠）。这里：
    1) 去掉 CJK 字符之间的单个空格；2) 收敛行尾空白与多余空行。
    """
    import re

    text = re.sub(r"(?<=[\u4e00-\u9fff]) +(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _looks_docx(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return "word/document.xml" in zf.namelist()
    except Exception:
        return False


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # 单页失败不致命
            logger.warning("PDF 单页解析失败: %s", exc)
    return "\n".join(pages).strip()


def _parse_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:  # 表格按行拼接，尽量保全文本
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "gbk", "gb18030"):
        try:
            return data.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


def split_text(text: str, chunk_size: int = 500, chunk_overlap: int = 50) -> list[str]:
    """中文优先的递归切分：段落 > 换行 > 中文句读 > 逗号 > 空格 > 字符。"""
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=max(50, int(chunk_size)),
        chunk_overlap=max(0, int(chunk_overlap)),
        separators=["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""],
        length_function=len,
    )
    return [c.strip() for c in splitter.split_text(text or "") if c.strip()]
